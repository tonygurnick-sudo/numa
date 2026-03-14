"""Native XLSX extraction using openpyxl."""

import structlog
from openpyxl import load_workbook

logger = structlog.get_logger(__name__)


def extract_xlsx(file_path: str) -> list[dict]:
    """Extract data from an XLSX file. Each sheet becomes a page."""
    try:
        wb_formulas = load_workbook(file_path, data_only=False, read_only=True)
        wb_values = load_workbook(file_path, data_only=True, read_only=True)
    except Exception as e:
        logger.error("Failed to open XLSX", error=str(e))
        return [
            {"page_number": 1, "num_words": 0, "text": f"(Failed to open XLSX: {e})"}
        ]

    pages: list[dict] = []

    for idx, sheet_name in enumerate(wb_formulas.sheetnames):
        try:
            ws_formula = wb_formulas[sheet_name]
            ws_value = wb_values[sheet_name]

            lines: list[str] = [f"# Sheet: {sheet_name}", ""]

            # Read all rows
            formula_rows = list(ws_formula.iter_rows(values_only=True))
            value_rows = list(ws_value.iter_rows(values_only=True))

            if not formula_rows:
                lines.append("(Empty sheet)")
            else:
                # Header row
                if formula_rows:
                    headers = [str(c) if c is not None else "" for c in formula_rows[0]]
                    lines.append("\t".join(headers))

                # Data rows
                for row_idx in range(1, len(formula_rows)):
                    cells = []
                    for col_idx, cell in enumerate(formula_rows[row_idx]):
                        if cell is None:
                            cells.append("")
                        elif isinstance(cell, str) and cell.startswith("="):
                            # Formula — show computed value too
                            computed = ""
                            if row_idx < len(value_rows) and col_idx < len(
                                value_rows[row_idx]
                            ):
                                computed = value_rows[row_idx][col_idx]
                            cells.append(
                                f"{cell} (={computed})"
                                if computed is not None
                                else str(cell)
                            )
                        else:
                            cells.append(str(cell))
                    lines.append("\t".join(cells))

            text = "\n".join(lines)
            pages.append(
                {
                    "page_number": idx + 1,
                    "num_words": len(text.split()),
                    "text": text,
                }
            )
        except Exception as e:
            logger.warning("Failed to extract sheet", sheet=sheet_name, error=str(e))
            pages.append(
                {
                    "page_number": idx + 1,
                    "num_words": 0,
                    "text": f"(Failed to extract sheet '{sheet_name}': {e})",
                }
            )

    wb_formulas.close()
    wb_values.close()

    if not pages:
        pages = [{"page_number": 1, "num_words": 0, "text": "(Empty XLSX workbook)"}]

    logger.info("XLSX extraction complete", sheets=len(pages))
    return pages
