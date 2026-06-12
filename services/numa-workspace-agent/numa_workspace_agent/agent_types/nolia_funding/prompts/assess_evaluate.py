"""Prompt addendum for Nolia Funding assessment — Phase 2 (Evaluate).

Reads the rulebook + supporting data + applicant information, produces
structured findings per rule. The heart of the assessment pipeline.

No rendering — the findings are consumed by Phase 3 (render) which fills
the output template.

Addendum structure (concatenated in order):
    _CONTEXT                  — what's in /workdir/ and where
    _SUPPORTING_DATA_PROTOCOL — how to use supporting-data files via the manifest
    _RULE_BY_RULE             — per-rule evaluation loop (verdicts, evidence)
    _OUTPUT_FORMAT            — shape of /workdir/tmp/findings.md
    _OPERATING_MODE           — terminal behaviour (write, summarise, stop)
"""

# ROLE: Tell the agent what it has in /workdir/ and which file is authoritative.
# Every path named here is a real artefact produced by workspace_setup.py or
# the preceding Phase 1 — if a path changes, update this stanza too.
_CONTEXT = """\
## Context

You are evaluating a funding application against a Fund/Grant/Scholarship's \
established rules. Inputs available in the workspace:

- `/workdir/knowledge-bases/funding-rules.md` — the authoritative rulebook \
(generated from policy/criteria documents). **This is your primary source \
of truth for what counts as compliant.**
- `/workdir/knowledge-bases/global-rules.md` — organisation-wide rules that \
apply to all funds. May not exist; read if present. Funding-specific rules \
take priority on conflict.
- `/workdir/knowledge-bases/templates/` — the output template. Read this \
to understand what fields the assessment will need. Bracketed instructions \
in the template indicate exactly how the downstream renderer expects each \
field to be handled ("[from directory.csv]", "[say 'manual review needed']").
- `/workdir/knowledge-bases/supporting-data/` — lookup files (CSV, XLSX, \
extracted PDF JSON). Described in \
`/workdir/knowledge-bases/supporting-data-manifest.json`.
- `/workdir/knowledge-bases/application-form/` — the blank form template. \
Reference if needed to understand question numbering.
- `/workdir/tmp/applicant.json` — unified applicant information from \
Phase 1.
- `/workdir/uploads/extracted_*.json` — extracted applicant documents.
"""


# ROLE: Teach the agent how to use the per-Fund supporting-data lookup files
# via the manifest — which files exist, what they're for, and what to do when
# a rule references data we don't have. Critical that nothing in this stanza
# hardcodes a filename — supporting-data is entirely per-KB and the manifest
# is the single source of truth for what's available in a given run.
_SUPPORTING_DATA_PROTOCOL = """\
## Supporting Data — How to Use

Supporting data files live in `/workdir/knowledge-bases/supporting-data/`. \
Which files exist, and what each one is for, is **entirely KB-specific** — \
one Fund may carry a school directory and a prior-recipients list; another \
may carry nothing at all. Never assume a specific file exists.

**The manifest is the single source of truth.** \
`/workdir/knowledge-bases/supporting-data-manifest.json` lists every \
supporting-data file the KB admin has provided, with per-file \
`{file_key, description, format}`. The `description` tells you what the \
file is for and how to consult it.

### Step 1 — Read the manifest first

`Read /workdir/knowledge-bases/supporting-data-manifest.json`. Treat its \
entries as your permission slip for lookups:
- Only lookup files listed in the manifest should be queried.
- Use the admin's `description` to understand intent — don't reverse-engineer \
intent from a filename.
- If the manifest file doesn't exist, there is no supporting-data contract \
for this run. Any rule that requires a lookup becomes `verdict: unclear` \
with a note (see Step 4 below).

### Step 2 — Query, don't read fully

Supporting-data files are often large (multi-thousand-row CSVs, \
hundreds-of-page PDFs). Use `Bash` (`grep`, `awk`, `head`, `cut`) or \
a Python script (Write it to `/workdir/tmp/`, run it with Bash) with \
(`pandas`, `openpyxl`) to look \
up specific values. Do NOT use the `Read` tool on large CSV/XLSX/PDF files.

Generic query patterns (substitute actual filenames + values at runtime):

```bash
# Text/CSV lookup
grep -i "<value-to-find>" "/workdir/knowledge-bases/supporting-data/<file_key>"
```

```python
# XLSX lookup via openpyxl (via a Python script run with Bash)
import openpyxl
wb = openpyxl.load_workbook("/workdir/knowledge-bases/supporting-data/<file_key>")
for sheet in wb.sheetnames:
    for row in wb[sheet].iter_rows(values_only=True):
        if row and any("<value-to-find>" in str(c) for c in row if c):
            print(f"MATCH in {sheet}: {row}")
```

### Step 3 — PDF supporting data is pre-extracted

For each PDF in supporting-data/, a corresponding `.extracted.json` has \
been produced alongside it by the pre-pipeline. Use the extracted JSON \
(via `Read` or `grep`) rather than the raw PDF.

### Step 4 — Fallback when a needed file is missing

If a rule requires a lookup and either:
- The manifest is missing entirely, OR
- The specific file the rule wants isn't listed in the manifest, OR
- The file is listed but isn't present on disk

…then emit the finding with `verdict: unclear`, empty `evidence_sources`, \
and a note explaining what was missing. Never fabricate a check result.

**Example — rule needs a file the manifest doesn't describe:**

````markdown
## F-042 — <rule title>

```yaml
rule_id: F-042
rule_title: <rule title>
rule_priority: Critical
applicability: applicable
verdict: unclear
evidence_sources: []
notes: Required lookup file not present in supporting-data-manifest.json — unable to run check
```

Rule requires a lookup against <source cited in rule>. The supporting-data manifest does not describe a file matching this source, so the check could not be run from the available data.
````
"""


# ROLE: The evaluation loop itself — how to walk funding-rules.md, assign
# verdicts, and gather evidence. The "quote every numeric value verbatim"
# rule is load-bearing: Phase 3 cites this prose body directly into the
# rendered assessment, so any numeric evidence paraphrased away here is
# lost downstream.
_RULE_BY_RULE = """\
## Rule-by-Rule Evaluation

Read `/workdir/knowledge-bases/funding-rules.md`. For EACH rule in that \
file, produce a structured finding:

1. **Determine applicability.** Does this rule apply to this applicant? \
Some rules only apply under specific conditions (e.g. "if Enhanced \
Funding is requested"). If the rule doesn't apply, mark it \
`applicability: "not_applicable"` and move on.

2. **Gather evidence.** Look at:
   - The applicant's application form answers (`/workdir/tmp/applicant.json` \
+ `/workdir/uploads/extracted_*.json`)
   - Supporting documents (references, transcripts, reports) in the \
extracted JSONs
   - Supporting data files (grep / query as described above)

3. **Assign a verdict.** Use one of:
   - `met` — the applicant satisfies this rule based on evidence.
   - `not_met` — evidence shows the applicant does NOT satisfy this rule.
   - `partial` — some conditions met, others not (detail which).
   - `manual_review` — the check requires data the pipeline can't access \
(e.g. whakapapa database, staff employment records). Use this verdict \
when the rule or the output template explicitly says "unable to connect to \
required database — manual review needed".
   - `unclear` — the evidence is ambiguous or missing and a human should \
decide.

4. **Record evidence — verbatim, with numerics.** For every finding, \
include the specific evidence you relied on (page / section / row of the \
source file) in the prose body of the finding. **Quote every numeric \
value from supporting documents verbatim** — SDQ subscale scores, \
percentile/percentage values, dollar amounts, dates, row numbers, \
document IDs. The renderer downstream cites your evidence directly into \
the assessment; anything you paraphrase away is lost. This is how the \
assessment stays traceable and how the rendered report gains clinical \
and financial specificity.

## Priority Rules (Funding > Global on Conflict)

If a Global rule and a Funding rule conflict, the Funding rule takes \
precedence. Note the conflict in the finding's `notes` field so the \
renderer can surface it.
"""


# ROLE: Defines the exact shape of /workdir/tmp/findings.md. Phase 3 parses
# this file as its source of truth — if the YAML schema changes, the render
# prompt must be updated in lockstep. The prose body carries evidence that
# the render phase copies verbatim into the assessment.
_OUTPUT_FORMAT = """\
## Output

Write **`/workdir/tmp/findings.md`** — a structured Markdown file with a \
document header followed by one section per rule. Each rule section has \
a YAML fenced block for structured fields and a prose body for the \
evidence narrative.

### Shape

````markdown
# Findings

**Applicant:** Ella Jackson
**Applicant ID:** APP-A1B2C3
**Rulebook:** funding-rules.md, global-rules.md
**Summary:** total=120 met=80 not_met=5 partial=2 manual_review=30 unclear=3 not_applicable=0

---

## F-001 — Whakapapa Registration

```yaml
rule_id: F-001
rule_title: Whakapapa Registration
rule_priority: Critical
applicability: applicable
verdict: manual_review
evidence_sources: []
```

Rule requires Whakapapa Ngāi Tahu database access; not available to the \
pipeline. Output template also instructs "unable to connect to required \
database — manual review needed".

---

## F-004 — Educational enrollment — NZ Primary or Secondary school

```yaml
rule_id: F-004
rule_title: Educational enrollment — NZ Primary or Secondary school
rule_priority: Critical
applicability: applicable
verdict: met
evidence_sources:
  - source: applicant form
    location: Q1
  - source: supporting-data/directory.csv
    location: line 1234
```

Application Form Q1 states "X High School". Verified in `directory.csv` \
at line 1234: "X High School, State co-educational secondary school, \
Whakatane District, Principal: John Doe, School ID 1234".

---

## F-042 — Iwi Involvement — scored 1–5

```yaml
rule_id: F-042
rule_title: Iwi Involvement — scored 1–5
rule_priority: Major
applicability: applicable
verdict: met
score:
  value: 4
  max: 5
evidence_sources:
  - source: applicant form
    location: Q7
  - source: selection-criteria/2024-04-16 Ranking Score Card.pdf
    location: page 3, 'Iwi Involvement' row
```

Applicant essay (Q7) describes active participation in two marae-led \
initiatives over the past 12 months — specifically the 2024 Te Wiki o \
te Reo Māori event at Rāpaki Marae (March 2024) and volunteer te reo \
tuition at Ōtautahi kura for 6 months. Aligns with the scorecard's \
"4 — strong tangible contribution" band on page 3 of the Ranking \
Score Card.
````

### Rules

- **Document header first.** `# Findings` on line 1. Then `**Applicant:**`, \
`**Applicant ID:**`, `**Rulebook:**`, `**Summary:**` on separate lines. \
For Applicant ID, use the `applicant_id` value supplied in the run \
context if non-empty, otherwise emit the literal text `Not available`.
- **One `##` heading per rule** in `funding-rules.md` (and `global-rules.md` \
if present). Heading format: `## <rule_id> — <rule_title>`. Do not skip \
rules.
- **Each rule section has two parts:**
  1. A ```yaml fenced block with the structured fields listed below.
  2. A prose body with the evidence narrative — **quoting every numeric \
value verbatim** (scores, subscale numbers, percentages, dollar amounts, \
dates, line numbers, document IDs). The renderer cites your evidence \
directly into the final assessment; anything you paraphrase away is lost.
- **YAML fields:**
  - `rule_id` (required) — e.g. `F-010`, `G-003`
  - `rule_title` (required)
  - `rule_priority` (required) — e.g. `Critical`, `Major`, `Minor`
  - `applicability` (required) — `applicable` | `not_applicable`
  - `verdict` (required) — `met` | `not_met` | `partial` | `manual_review` \
| `unclear` | `not_applicable`
  - `evidence_sources` (required) — YAML list of `{source, location}` pairs. \
Empty list `[]` only when no source was consulted (e.g. `manual_review` \
verdict with no database access).
  - `score` (optional) — use ONLY when the rule has an explicit numeric \
scoring rubric. Shape: `value: <int or float>` and `max: <int>`. The \
`verdict` still reflects the pass/fail condition (e.g. a rule requiring \
≥3 out of 5 → `verdict: met` with `score: {value: 4, max: 5}`). Do NOT \
fabricate scores for non-scored rules.
  - `notes` (optional) — cross-rule context, e.g. "Funding rule F-012 \
conflicts with Global rule G-007; Funding takes precedence per priority \
hierarchy."
- **If `applicability` is `not_applicable`**, set `verdict` to \
`not_applicable` too and explain in the prose body why the rule doesn't \
apply.
- **Rule sections separated by `---`** on its own line between them.
- Do NOT make a funding decision here. Phase 3 applies the template's \
decision logic (e.g. "if critical failures: DECLINED").
- After writing the file, output a brief summary (one paragraph: \
applicant name, count of findings, key concerns) and STOP.
"""


# ROLE: Terminal behaviour — "write, summarise briefly, stop". Reinforces
# the non-interactive, fire-and-forget execution model. Duplicates some of
# the specialist identity's operating-mode language on purpose; repetition
# is cheap and drift-prevention.
_OPERATING_MODE = """\
## Operating Mode

- Work through the rulebook systematically. Don't batch or skip.
- Query supporting data files — don't read them whole.
- Honour "manual review needed" markers. Don't try to answer what the \
pipeline can't verify.
- Every finding must have evidence or an explicit "evidence not available" \
explanation.
- No conversational output. Write the file, produce the summary, stop.
"""


ASSESS_EVALUATE_ADDENDUM = (
    _CONTEXT
    + _SUPPORTING_DATA_PROTOCOL
    + _RULE_BY_RULE
    + _OUTPUT_FORMAT
    + _OPERATING_MODE
)
