# pylint: disable=protected-access
import dataclasses
import io
import os
import sys
import unittest
import unittest.mock
from typing import List
from unittest.mock import MagicMock, patch

import lambda_function  # pylint: disable=wrong-import-position

# Add the lib directory to the Python path to find the modules
project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../../"))
sys.path.append(os.path.join(project_root, "lib/aws-transcribe"))
sys.path.append(os.path.join(project_root, "lib/helpers"))

# Mock JWT module
sys.modules["jwt"] = unittest.mock.Mock()  # type: ignore
sys.modules["aws_transcribe"] = unittest.mock.Mock()  # type: ignore

# Create proper mock for fm_vision_extraction with the classes we need
haiku_mock = unittest.mock.Mock()

# Create proper dataclasses for inheritance


@dataclasses.dataclass
class MockDocumentPage:
    page_number: int = 0
    num_words: int = 0
    text: str = ""


@dataclasses.dataclass
class MockDocument:
    name: str = ""
    num_pages: int = 0
    total_num_words: int = 0
    pages: List[MockDocumentPage] = dataclasses.field(default_factory=list)


haiku_mock.DocumentPage = MockDocumentPage
haiku_mock.Document = MockDocument
haiku_mock.extract_content = unittest.mock.Mock()

sys.modules["fm_vision_extraction"] = haiku_mock

# Import docx modules after mocks are set up - they're conditionally available in tests
try:
    import docx  # noqa: F401 # pylint: disable=unused-import
    from docx import Document as DocxDoc
    from docx.enum.text import WD_BREAK

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    DocxDoc = object  # type: ignore[misc, assignment]
    WD_BREAK = object()  # type: ignore[misc, assignment]


class TestException(Exception):
    pass


class TestLambdaFunction(unittest.TestCase):
    @patch("lambda_function.s3_client")
    def test_handler_txt_file(self, mock_s3_client):
        """Test processing of text files"""
        mock_s3_client.get_object.return_value = {
            "Body": MagicMock(read=lambda: b"Sample text content")
        }
        event = {
            "input_bucket": "test-bucket",
            "input_key": "test.txt",
            "return_content": True,
        }

        response = lambda_function.handler(event, {})

        self.assertEqual(response["content"], "Sample text content\n")
        mock_s3_client.get_object.assert_called_once_with(
            Bucket="test-bucket",
            Key="test.txt",
        )
        # Should be called 3 times: initial status, main output, final status
        self.assertEqual(mock_s3_client.put_object.call_count, 3)
        kwargs = mock_s3_client.put_object.call_args.kwargs
        self.assertEqual(kwargs["Bucket"], "test-bucket")
        self.assertEqual(kwargs["Key"], "test.txt.status.json")

    @patch("lambda_function.s3_client")
    def test_handler_empty_txt_file(self, mock_s3_client):
        """Test processing of empty text files"""
        mock_s3_client.get_object.return_value = {
            "Body": MagicMock(read=lambda: b""),
        }
        event = {
            "input_bucket": "test-bucket",
            "input_key": "empty.txt",
            "return_content": True,
        }

        response = lambda_function.handler(event, {})

        self.assertEqual(response["content"], "\n")
        mock_s3_client.get_object.assert_called_once_with(
            Bucket="test-bucket",
            Key="empty.txt",
        )

    @patch("lambda_function.s3_client")
    def test_handler_unsupported_filetype_error(self, mock_s3_client):
        """Test handling of unsupported file types"""
        event = {
            "input_bucket": "test-bucket",
            "input_key": "unsupported.xyz",
        }

        with self.assertRaises(lambda_function.UnsupportedFileFormat):
            lambda_function.handler(event, {})

        mock_s3_client.get_object.assert_not_called()
        # Should be called 2 times: initial status, final error status
        self.assertEqual(mock_s3_client.put_object.call_count, 2)

    @patch("lambda_function.s3_client")
    def test_handler_s3_exception(self, mock_s3_client):
        """Test handling of S3 exceptions"""
        mock_s3_client.get_object.side_effect = TestException("S3 error")
        event = {
            "input_bucket": "test-bucket",
            "input_key": "error.txt",
        }

        with self.assertRaises(TestException):
            lambda_function.handler(event, {})

        mock_s3_client.get_object.assert_called_once_with(
            Bucket="test-bucket",
            Key="error.txt",
        )

    @patch("lambda_function.fm_vision_extraction")
    @patch("lambda_function.s3_client")
    def test_handler_haiku_supported_file(self, mock_s3_client, mock_haiku):
        """Test processing of Haiku 3 supported files (PDF, DOCX, images)"""
        mock_document = lambda_function.Document(
            name="test.pdf",
            num_pages=1,
            pages=[
                lambda_function.DocumentPage(
                    page_number=1, num_words=3, text="Sample PDF content"
                )
            ],
            total_num_words=3,
        )
        mock_haiku.extract_content.return_value = mock_document

        event = {
            "input_bucket": "test-bucket",
            "input_key": "test.pdf",
            "return_content": True,
        }

        response = lambda_function.handler(event, {})

        self.assertEqual(response["content"], "Sample PDF content\n")
        mock_haiku.extract_content.assert_called_once_with(
            "test-bucket", "test.pdf", None
        )
        # Should be called 3 times: initial status, main output, final status
        self.assertEqual(mock_s3_client.put_object.call_count, 3)

    @patch("lambda_function.s3_client")
    def test_handler_excel_file(self, mock_s3_client):
        """Test processing of Excel files"""
        # Mock Excel file content
        mock_excel_content = b"fake excel content"
        mock_s3_client.get_object.return_value = {
            "Body": MagicMock(read=lambda: mock_excel_content)
        }

        # Mock the extract_excel_structure_simple function (default behavior)
        with patch("lambda_function.extract_excel_structure_simple") as mock_extract:
            mock_extract.return_value = {
                1: {
                    "sheet_name": "Sheet1",
                    "structure": {
                        "columns": ["A", "B"],
                        "headers": ["A1: Header", "B1: Value"],
                        "rows_count": 1,
                    },
                    "rows": ["A1: Header | B1: Value"],
                }
            }

            event = {
                "input_bucket": "test-bucket",
                "input_key": "test.xlsx",
                "return_content": True,
            }

            response = lambda_function.handler(event, {})

            self.assertEqual(response["content"], "A1: Header | B1: Value\n")
            mock_s3_client.get_object.assert_called_once_with(
                Bucket="test-bucket", Key="test.xlsx"
            )
            # Should be called 3 times: initial status, main output, final status
            self.assertEqual(mock_s3_client.put_object.call_count, 3)

    def test_extract_docx_pages_includes_table_text(self):
        """DOCX extraction should capture table content."""
        if not DOCX_AVAILABLE:
            self.skipTest("python-docx not available")

        buffer = io.BytesIO()
        doc = DocxDoc()  # type: ignore[misc]
        doc.add_paragraph("Intro paragraph")  # type: ignore[attr-defined]
        table = doc.add_table(rows=2, cols=2)  # type: ignore[attr-defined]
        table.cell(0, 0).text = "Label"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Notes"
        table.cell(1, 1).text = "More detail"
        doc.add_paragraph("Closing paragraph")  # type: ignore[attr-defined]
        doc.save(buffer)  # type: ignore[attr-defined]

        pages = lambda_function.extract_docx_pages(buffer.getvalue())

        self.assertIn("Intro paragraph", pages[1])
        self.assertIn("Label | Value", pages[1])
        self.assertIn("Notes | More detail", pages[1])
        self.assertIn("Closing paragraph", pages[1])

    def test_extract_docx_pages_splits_table_on_page_break(self):
        """A page break inside a table cell should start a new page."""
        if not DOCX_AVAILABLE:
            self.skipTest("python-docx not available")

        buffer = io.BytesIO()
        doc = DocxDoc()  # type: ignore[misc]
        table = doc.add_table(rows=1, cols=1)  # type: ignore[attr-defined]
        paragraph = table.cell(0, 0).paragraphs[0]
        paragraph.add_run("First page content")
        paragraph.add_run().add_break(WD_BREAK.PAGE)  # type: ignore[attr-defined]
        paragraph.add_run("Second page content")
        doc.save(buffer)  # type: ignore[attr-defined]

        pages = lambda_function.extract_docx_pages(buffer.getvalue())

        self.assertIn("First page content", pages[1])
        self.assertNotIn("Second page content", pages[1])
        self.assertIn("Second page content", pages[2])

    @patch("lambda_function.aws_transcribe")
    @patch("lambda_function.s3_client")
    def test_handler_audio_file(self, mock_s3_client, mock_transcribe):
        """Test processing of audio files"""
        mock_response = MagicMock()
        mock_response.text = "Sample audio transcription"
        mock_transcribe.transcribe.return_value = mock_response

        event = {
            "input_bucket": "test-bucket",
            "input_key": "test.mp3",
            "return_content": True,
        }

        response = lambda_function.handler(event, {})

        self.assertEqual(response["content"], "Sample audio transcription\n")
        mock_transcribe.transcribe.assert_called_once()
        # Should be called 3 times: initial status, main output, final status
        self.assertEqual(mock_s3_client.put_object.call_count, 3)

    def test_text_to_document(self):
        """Test text to document conversion"""
        document = lambda_function._text_to_document(
            "Hello world test", "test.txt", None
        )

        self.assertEqual(document.name, "test.txt")
        self.assertEqual(document.num_pages, 1)
        self.assertEqual(document.total_num_words, 3)
        self.assertEqual(len(document.pages), 1)
        self.assertEqual(document.pages[0].text, "Hello world test")
        self.assertEqual(document.pages[0].num_words, 3)
        self.assertEqual(document.pages[0].page_number, 1)

    def test_document_to_string(self):
        """Test document to string conversion"""
        document = lambda_function.Document(
            name="test.txt",
            num_pages=2,
            pages=[
                lambda_function.DocumentPage(
                    page_number=1, num_words=2, text="First page"
                ),
                lambda_function.DocumentPage(
                    page_number=2, num_words=2, text="Second page"
                ),
            ],
            total_num_words=4,
        )

        result = lambda_function._document_to_string(document)
        self.assertEqual(result, "First page\nSecond page\n")


if __name__ == "__main__":
    unittest.main()
