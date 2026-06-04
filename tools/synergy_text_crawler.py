#!/usr/bin/env python3
"""Crawl ALL text from a 12d Synergy instance into a Bedrock-KB-ready corpus.

Why this exists: 12d has no cross-project search. To answer questions like
"which jobs use material X", you pull the *text* out of Synergy's documents and
index it yourself. You do NOT copy the 8 TB — that's mostly CAD/binary geometry.
This script downloads only text-bearing documents (PDF/Word/Excel/etc.), extracts
their text, and writes one `.txt` + one `.metadata.json` sidecar per document in
the exact layout Bedrock Knowledge Base ingests from S3 (metadata lets you
filter/attribute hits back to a job — that's your cross-job search).

Enumeration: jobs are searched globally (`/jobs/search`), then files per job via
`/files/search` scoped to the job + sub-jobs (`LimitSearchTo:2`, the only working
file-search mode — there is no global file search). Files are deduped by id, so
overlap from sub-jobs is harmless.

Resumable: every processed file id is recorded in `manifest.jsonl`; re-running
skips them, so you can crawl 13k jobs in stages or recover from interruptions.

Usage:
    export SYNERGY_PAT='<your 12d personal access token>'
    python3 synergy_text_crawler.py --out ./synergy_corpus

    # or pull the PAT straight from the cuttriss vault (needs AWS_PROFILE=cuttriss):
    python3 synergy_text_crawler.py --from-vault \
        cuttriss/vault/users/d99ee478-8021-709d-67fb-d246feea7bed \
        --vault-region ap-southeast-2 --out ./synergy_corpus

    # smoke test on a few jobs first:
    python3 synergy_text_crawler.py --out ./test --max-jobs 5

Then sync the output to S3 and point a Bedrock KB (S3 data source) at it:
    aws s3 sync ./synergy_corpus/docs s3://<your-kb-bucket>/synergy/

Dependencies: httpx (required). Text extractors are optional and used when
present: pdfplumber or PyPDF2 (PDF), python-docx (.docx), openpyxl (.xlsx),
python-pptx (.pptx). Missing extractor => that file type is skipped with a count.
    pip install httpx pdfplumber python-docx openpyxl python-pptx
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Tuple

import httpx

# --------------------------------------------------------------------------- #
# File-type policy
# --------------------------------------------------------------------------- #
# Only these are downloaded + text-extracted. Everything else (CAD, models,
# point clouds, images, archives) is skipped — that's where the terabytes are
# and there are no words in them.
TEXT_EXTS = {
    "pdf",
    "docx",
    "xlsx",
    "xlsm",
    "pptx",
    "txt",
    "csv",
    "tsv",
    "md",
    "rtf",
    "xml",
    "json",
    "htm",
    "html",
    "log",
}
# Legacy binary Office (need LibreOffice/textract) — skipped but counted so you
# know how much you're leaving on the table.
LEGACY_EXTS = {"doc", "xls", "ppt"}


# --------------------------------------------------------------------------- #
# Synergy API client
# --------------------------------------------------------------------------- #
class Synergy:
    def __init__(self, base_url: str, pat: str, timeout: float = 120.0):
        self.base = base_url.rstrip("/")
        self.h = {
            "Authorization": f"Bearer {pat.strip()}",
            "Content-Type": "application/json",
        }
        self.client = httpx.Client(timeout=timeout)

    def _post(self, path: str, body: dict, retries: int = 4) -> httpx.Response:
        url = f"{self.base}{path}"
        for attempt in range(retries):
            try:
                r = self.client.post(url, headers=self.h, json=body)
            except httpx.HTTPError as exc:
                if attempt == retries - 1:
                    raise
                time.sleep(2**attempt)
                continue
            if r.status_code in (401, 403):
                raise SystemExit(
                    f"\nAuth failed (HTTP {r.status_code}). Your PAT is expired or "
                    f"lacks permission. Generate a new one in Synergy and re-run.\n"
                )
            if r.status_code == 429 or r.status_code >= 500:
                if attempt == retries - 1:
                    return r
                time.sleep(2**attempt)
                continue
            return r
        return r

    def iter_jobs(self, page_size: int = 100) -> Iterator[dict]:
        """Yield every job (TopLevel=false returns all jobs, incl. sub-jobs)."""
        page = 1
        while True:
            body = {
                "Page": page,
                "PageSize": page_size,
                "Name": "",
                "QuickSearchTerm": "",
                "Attributes": [
                    {
                        "Attribute": {
                            "Name": "TopLevel",
                            "DisplayName": "Restrict to top level?",
                        },
                        "Type": "SynergyServerWeb.API.Models.SelectableProgrammaticAttribute",
                        "Value": False,
                        "SearchQueryType": 4,
                        "Operation": 0,
                        "Name": "Restrict to top level?",
                        "OperationName": "=",
                    }
                ],
            }
            r = self._post("/api/v1/jobs/search", body)
            r.raise_for_status()
            d = r.json()
            rows = d.get("Result") or d.get("Items") or []
            for row in rows:
                yield row
            total_pages = d.get("TotalPages") or 0
            if page >= total_pages or not rows:
                break
            page += 1

    def iter_job_files(self, job_idstring: str, page_size: int = 100) -> Iterator[dict]:
        """Yield every file in a job + its sub-jobs (empty filter = match all)."""
        parts = job_idstring.split("_")
        if len(parts) != 2 or not all(p.isdigit() for p in parts):
            return
        limit_id = {
            "IDString": job_idstring,
            "_id": int(parts[0]),
            "_server_id": int(parts[1]),
        }
        page = 1
        while True:
            body = {
                "Page": page,
                "PageSize": page_size,
                "Attributes": [],
                "ShowDeletedFiles": False,
                "FileName": "",
                "Contents": "",
                "LimitSearchTo": 2,  # job + sub-jobs — the only working file scope
                "LimitID": limit_id,
            }
            r = self._post("/api/v1/files/search", body)
            if r.status_code >= 400:
                # Don't let one bad job kill the whole crawl.
                sys.stderr.write(
                    f"  ! files/search failed for job {job_idstring} "
                    f"(HTTP {r.status_code}); skipping job\n"
                )
                return
            d = r.json()
            rows = d.get("Result") or d.get("Items") or []
            for row in rows:
                yield row
            total_pages = d.get("TotalPages") or 0
            if page >= total_pages or not rows:
                break
            page += 1

    def download(self, file_id: str, version: int) -> bytes:
        """Download raw file bytes (POST with empty body — verified working)."""
        url = (
            f"{self.base}/api/v1/files/{file_id}/download"
            f"?version={version}&with_references=false"
        )
        r = self.client.post(
            url,
            headers={**self.h, "Content-Type": "application/octet-stream"},
            content=b"",
        )
        if r.status_code in (401, 403):
            raise SystemExit(f"Auth failed on download (HTTP {r.status_code}).")
        r.raise_for_status()
        return r.content


# --------------------------------------------------------------------------- #
# Text extraction (each extractor is optional)
# --------------------------------------------------------------------------- #
def extract_text(data: bytes, ext: str) -> str:
    if ext in ("txt", "csv", "tsv", "md", "log", "xml", "json"):
        return data.decode("utf-8", errors="replace")
    if ext in ("htm", "html"):
        try:
            from bs4 import BeautifulSoup

            return BeautifulSoup(data, "html.parser").get_text(" ", strip=True)
        except ImportError:
            return re.sub(r"<[^>]+>", " ", data.decode("utf-8", errors="replace"))
    if ext == "pdf":
        return _pdf(data)
    if ext == "docx":
        return _docx(data)
    if ext in ("xlsx", "xlsm"):
        return _xlsx(data)
    if ext == "pptx":
        return _pptx(data)
    if ext == "rtf":
        # crude RTF strip; good enough for indexing
        txt = data.decode("latin-1", errors="replace")
        txt = re.sub(r"\\'[0-9a-fA-F]{2}", " ", txt)
        txt = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", txt)
        return re.sub(r"[{}]", " ", txt)
    raise _SkipType(ext)


def _pdf(data: bytes) -> str:
    import io

    try:
        import pdfplumber

        out = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                out.append(page.extract_text() or "")
        return "\n".join(out)
    except ImportError:
        pass
    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    except ImportError:
        raise _SkipType("pdf (install pdfplumber or PyPDF2)")


def _docx(data: bytes) -> str:
    import io

    try:
        import docx
    except ImportError:
        raise _SkipType("docx (install python-docx)")
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    import io

    try:
        import openpyxl
    except ImportError:
        raise _SkipType("xlsx (install openpyxl)")
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append("\t".join(cells))
    return "\n".join(out)


def _pptx(data: bytes) -> str:
    import io

    try:
        from pptx import Presentation
    except ImportError:
        raise _SkipType("pptx (install python-pptx)")
    prs = Presentation(io.BytesIO(data))
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n".join(out)


class _SkipType(Exception):
    pass


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def file_ext(name: str) -> str:
    name = (name or "").lower()
    return name.rsplit(".", 1)[-1] if "." in name else ""


def safe(name: str, limit: int = 120) -> str:
    name = re.sub(r"[^\w.\-]+", "_", name or "").strip("_.")
    return (name or "file")[:limit]


def get(d: dict, *keys, default=None):
    for k in keys:
        v = d.get(k)
        if v not in (None, ""):
            return v
    return default


def id_string(obj: dict) -> Optional[str]:
    return (obj.get("ID") or {}).get("IDString") or obj.get("IDString")


def attr_map(file_obj: dict) -> Dict[str, str]:
    """Flatten the Attributes array into {display_name: value}."""
    out: Dict[str, str] = {}
    for a in file_obj.get("Attributes") or []:
        if not isinstance(a, dict):
            continue
        key = a.get("DisplayName") or a.get("Name")
        val = a.get("value")
        if isinstance(val, dict):
            val = val.get("_value")
        if key and val not in (None, ""):
            out[str(key)] = str(val)
    return out


# --------------------------------------------------------------------------- #
# Main crawl
# --------------------------------------------------------------------------- #
def main() -> None:
    ap = argparse.ArgumentParser(description="Crawl Synergy text → Bedrock-KB corpus")
    ap.add_argument("--base-url", default="https://synergy.cuttriss.co.nz")
    ap.add_argument("--pat", default=os.environ.get("SYNERGY_PAT", ""))
    ap.add_argument(
        "--from-vault",
        metavar="SECRET_ID",
        help="Fetch PAT from AWS Secrets Manager (uses current AWS_PROFILE)",
    )
    ap.add_argument("--vault-region", default="ap-southeast-2")
    ap.add_argument("--out", default="./synergy_corpus")
    ap.add_argument("--page-size", type=int, default=100)
    ap.add_argument(
        "--max-mb",
        type=float,
        default=75.0,
        help="Skip text docs larger than this (MB) — guards against giant PDFs",
    )
    ap.add_argument("--max-jobs", type=int, default=0, help="Stop after N jobs (test)")
    args = ap.parse_args()

    pat = args.pat
    if args.from_vault:
        import boto3

        sm = boto3.client("secretsmanager", region_name=args.vault_region)
        sec = json.loads(sm.get_secret_value(SecretId=args.from_vault)["SecretString"])
        fields = (sec.get("secrets", {}).get("connector-synergy") or {}).get(
            "fields"
        ) or {}
        pat = (fields.get("access_token") or "").strip()
    if not pat:
        ap.error("No PAT. Set SYNERGY_PAT, pass --pat, or use --from-vault.")

    out_root = Path(args.out)
    docs_dir = out_root / "docs"
    docs_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = out_root / "manifest.jsonl"

    # Resume: skip file ids already written.
    done: set[str] = set()
    if manifest_path.exists():
        with manifest_path.open() as fh:
            for line in fh:
                try:
                    done.add(json.loads(line)["file_id"])
                except Exception:
                    pass
    print(f"Resuming: {len(done)} files already processed" if done else "Fresh crawl")

    syn = Synergy(args.base_url, pat)
    manifest = manifest_path.open("a")

    stats = {
        "jobs": 0,
        "files_seen": 0,
        "extracted": 0,
        "skipped_type": 0,
        "skipped_legacy": 0,
        "skipped_big": 0,
        "skipped_empty": 0,
        "errors": 0,
        "bytes_downloaded": 0,
    }
    seen_files: set[str] = set(done)
    max_bytes = int(args.max_mb * 1024 * 1024)
    t0 = time.time()

    try:
        for job in syn.iter_jobs(args.page_size):
            job_id = id_string(job)
            if not job_id:
                continue
            stats["jobs"] += 1
            job_name = get(job, "Name", default=job_id)
            job_path = get(job, "Path", default="")

            for f in syn.iter_job_files(job_id, args.page_size):
                fid = id_string(f)
                if not fid or fid in seen_files:
                    continue
                seen_files.add(fid)
                stats["files_seen"] += 1
                fname = get(f, "FileName", "Name", default="")
                ext = file_ext(fname)
                if ext in LEGACY_EXTS:
                    stats["skipped_legacy"] += 1
                    continue
                if ext not in TEXT_EXTS:
                    stats["skipped_type"] += 1
                    continue
                size = get(f, "FileSize", "Size", default=0) or 0
                try:
                    if int(size) > max_bytes:
                        stats["skipped_big"] += 1
                        continue
                except (TypeError, ValueError):
                    pass

                version = get(f, "LatestVersion", "latestVersion", default=1) or 1
                try:
                    data = syn.download(fid, version)
                    stats["bytes_downloaded"] += len(data)
                    text = extract_text(data, ext)
                except _SkipType:
                    stats["skipped_type"] += 1
                    continue
                except SystemExit:
                    raise
                except Exception as exc:  # noqa: BLE001
                    stats["errors"] += 1
                    sys.stderr.write(f"  ! {fid} {fname}: {exc}\n")
                    continue

                text = (text or "").strip()
                if not text:
                    stats["skipped_empty"] += 1
                    continue

                stem = f"{fid}__{safe(fname)}"
                (docs_dir / f"{stem}.txt").write_text(text, encoding="utf-8")
                meta = {
                    "metadataAttributes": {
                        "job_id": job_id,
                        "job_name": job_name,
                        "job_path": job_path,
                        "file_id": fid,
                        "file_name": fname,
                        "file_path": get(f, "Path", default=""),
                        "version": str(version),
                        **{f"attr_{safe(k)}": v for k, v in attr_map(f).items()},
                    }
                }
                (docs_dir / f"{stem}.txt.metadata.json").write_text(
                    json.dumps(meta, ensure_ascii=False), encoding="utf-8"
                )
                manifest.write(
                    json.dumps(
                        {
                            "file_id": fid,
                            "job_id": job_id,
                            "job_name": job_name,
                            "file_name": fname,
                            "chars": len(text),
                        }
                    )
                    + "\n"
                )
                manifest.flush()
                stats["extracted"] += 1

            if stats["jobs"] % 25 == 0:
                _progress(stats, t0)
            if args.max_jobs and stats["jobs"] >= args.max_jobs:
                break
    except KeyboardInterrupt:
        print("\nInterrupted — progress saved to manifest, safe to re-run.")
    finally:
        manifest.close()
        syn.client.close()

    _progress(stats, t0, final=True)
    print(f"\nCorpus written to: {docs_dir}")
    print("Next: aws s3 sync", str(docs_dir), "s3://<your-kb-bucket>/synergy/")


def _progress(stats: dict, t0: float, final: bool = False) -> None:
    mins = (time.time() - t0) / 60
    gb = stats["bytes_downloaded"] / (1024**3)
    tag = "DONE" if final else "..."
    print(
        f"[{tag}] jobs={stats['jobs']} files_seen={stats['files_seen']} "
        f"extracted={stats['extracted']} "
        f"skip(type={stats['skipped_type']},legacy={stats['skipped_legacy']},"
        f"big={stats['skipped_big']},empty={stats['skipped_empty']}) "
        f"err={stats['errors']} dl={gb:.2f}GB {mins:.1f}min",
        flush=True,
    )


if __name__ == "__main__":
    main()
