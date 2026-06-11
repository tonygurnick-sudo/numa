#!/usr/bin/env python3
"""Extract text from an already-mirrored Synergy backup in S3 — text layer only, zero LLM cost.

Sibling of `synergy_text_crawler.py`: that one pulls from the live Synergy API;
this one reads the files that are ALREADY in S3 (e.g. `numa-cuttriss-backup`) and
extracts only the embedded text layer. No Bedrock Vision, no OCR, no Transcribe —
the only spend is S3 GETs. Born-digital docs come out free; scanned PDFs and images
are skipped (and counted) so you can decide later whether they're worth a paid pass.

Writes one `.txt` + one `.txt.metadata.json` per document in the Bedrock-KB layout,
so you can `aws s3 sync` the output at a KB S3 data source.

Usage:
    AWS_PROFILE=cuttriss python3 synergy_s3_text_extractor.py \
        --bucket numa-cuttriss-backup --region us-east-2 \
        --prefix documents/ --out ./synergy_corpus

    # smoke test on one job:
    AWS_PROFILE=cuttriss python3 synergy_s3_text_extractor.py \
        --bucket numa-cuttriss-backup --region us-east-2 \
        --prefix documents/kb-626f1ef5-e63e-44c9-b35c-ad12c2166098/0812/ --out ./test

Deps: boto3 (required). Optional, used when present: pymupdf (PDF), python-docx
(.docx), openpyxl (.xlsx/.xlsm), extract-msg (.msg). Missing extractor => that
type is skipped with a count.
    pip install boto3 pymupdf python-docx openpyxl extract-msg
"""

from __future__ import annotations

import argparse
import email
import io
import json
import os
import re
import sys
import time
from email import policy as email_policy
from pathlib import Path
from typing import Iterator, Optional

import boto3

# Extensions we can extract for free. Everything else (CAD, models, point clouds,
# raw images, archives, audio/video) is skipped — no text layer, or only via paid
# vision/OCR which this tool deliberately does not do.
TEXT_DECODE = {
    "txt",
    "csv",
    "tsv",
    "md",
    "log",
    "xml",
    "json",
    "htm",
    "html",
    "yaml",
    "yml",
    "rtf",
}
PDF_EXTS = {"pdf"}
DOCX_EXTS = {"docx"}
XLSX_EXTS = {"xlsx", "xlsm"}
MSG_EXTS = {"msg"}
EML_EXTS = {"eml"}
HANDLED = TEXT_DECODE | PDF_EXTS | DOCX_EXTS | XLSX_EXTS | MSG_EXTS | EML_EXTS

# A PDF averaging fewer than this many stripped chars/page is treated as scanned
# (image-only) and skipped — extracting it would need vision/OCR.
MIN_PDF_CHARS_PER_PAGE = 25


def file_ext(name: str) -> str:
    name = (name or "").lower()
    return name.rsplit(".", 1)[-1] if "." in name else ""


def safe(name: str, limit: int = 150) -> str:
    name = re.sub(r"[^\w.\-]+", "_", name or "").strip("_.")
    return (name or "file")[:limit]


def job_id_from_key(key: str) -> str:
    """Best-effort job number: the path segment right after a `kb-...` prefix."""
    parts = key.split("/")
    for i, p in enumerate(parts):
        if p.startswith("kb-") and i + 1 < len(parts):
            return parts[i + 1]
    return ""


def html_to_text(s: str) -> str:
    s = re.sub(r"(?i)<\s*br\s*/?\s*>", "\n", s)
    s = re.sub(r"(?i)</\s*(p|div|li|tr|h[1-6])\s*>", "\n", s)
    s = re.sub(r"<[^>]+>", " ", s)
    return re.sub(r"[ \t]+", " ", s)


# --------------------------------------------------------------------------- #
# Per-type extractors. Each returns text or raises _Skip(reason).
# --------------------------------------------------------------------------- #
class _Skip(Exception):
    pass


def extract(data: bytes, ext: str) -> str:
    if ext in TEXT_DECODE:
        txt = data.decode("utf-8", errors="replace")
        return html_to_text(txt) if ext in ("htm", "html") else txt
    if ext in PDF_EXTS:
        return _pdf(data)
    if ext in DOCX_EXTS:
        return _docx(data)
    if ext in XLSX_EXTS:
        return _xlsx(data)
    if ext in MSG_EXTS:
        return _msg(data)
    if ext in EML_EXTS:
        return _eml(data)
    raise _Skip(f"unhandled:{ext}")


def _pdf(data: bytes) -> str:
    try:
        import fitz  # pymupdf
    except ImportError:
        raise _Skip("pdf (install pymupdf)")
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        pages = [(p.get_text("text") or "") for p in doc]
        n = doc.page_count or 1
    finally:
        doc.close()
    total = sum(len(p.strip()) for p in pages)
    if total / n < MIN_PDF_CHARS_PER_PAGE:
        raise _Skip("pdf-scanned")  # image-only — would need vision/OCR
    return "\n".join(p for p in pages if p.strip())


def _docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise _Skip("docx (install python-docx)")
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    try:
        import openpyxl
    except ImportError:
        raise _Skip("xlsx (install openpyxl)")
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append("\t".join(cells))
    return "\n".join(out)


def _msg(data: bytes) -> str:
    try:
        import extract_msg
    except ImportError:
        raise _Skip("msg (install extract-msg)")
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".msg", delete=False) as fh:
        fh.write(data)
        path = fh.name
    try:
        m = extract_msg.Message(path)
        body = m.body or ""
        if isinstance(body, bytes):
            body = body.decode("utf-8", errors="ignore")
        head = f"From: {m.sender}\nTo: {m.to}\nSubject: {m.subject}\nDate: {m.date}\n"
        if not body.strip() and m.htmlBody:
            hb = m.htmlBody
            body = html_to_text(
                hb.decode("utf-8", "ignore") if isinstance(hb, bytes) else hb
            )
        m.close()
        return head + "\n" + body
    finally:
        if os.path.exists(path):
            os.remove(path)


def _eml(data: bytes) -> str:
    msg = email.message_from_bytes(data, policy=email_policy.default)
    head = (
        f"From: {msg.get('From','')}\nTo: {msg.get('To','')}\n"
        f"Subject: {msg.get('Subject','')}\nDate: {msg.get('Date','')}\n"
    )
    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                body = part.get_content()
                break
        if not body:
            for part in msg.walk():
                if part.get_content_type() == "text/html":
                    body = html_to_text(part.get_content())
                    break
    else:
        body = msg.get_content()
        if msg.get_content_type() == "text/html":
            body = html_to_text(body)
    return head + "\n" + (body or "")


# --------------------------------------------------------------------------- #
def iter_keys(
    s3, bucket: str, prefix: str, start_after: str = ""
) -> Iterator[tuple[str, int]]:
    paginator = s3.get_paginator("list_objects_v2")
    kw = {"Bucket": bucket, "Prefix": prefix}
    if start_after:
        kw["StartAfter"] = start_after
    for page in paginator.paginate(**kw):
        for obj in page.get("Contents", []) or []:
            key = obj["Key"]
            if key.endswith(".metadata.json") or key.endswith("/"):
                continue
            yield key, obj.get("Size", 0)


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Extract text-layer only from an S3 Synergy backup"
    )
    ap.add_argument("--bucket", required=True)
    ap.add_argument("--region", default="us-east-2")
    ap.add_argument("--prefix", default="")
    ap.add_argument("--out", default="./synergy_corpus")
    ap.add_argument(
        "--max-mb", type=float, default=75.0, help="Skip files larger than this"
    )
    ap.add_argument(
        "--max-files", type=int, default=0, help="Stop after N handled files (test)"
    )
    ap.add_argument(
        "--tenant-id",
        default="",
        help="Numa CLIENT_NAME — REQUIRED for chat retrieval (filter: tenant_id==this)",
    )
    ap.add_argument(
        "--kb-id",
        default="",
        help="Numa folder kb_id — REQUIRED for chat retrieval (filter: kb_id==this)",
    )
    args = ap.parse_args()

    out_root = Path(args.out)
    docs_dir = out_root / "docs"
    docs_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = out_root / "manifest.jsonl"

    done: set[str] = set()
    if manifest_path.exists():
        with manifest_path.open() as fh:
            for line in fh:
                try:
                    done.add(json.loads(line)["key"])
                except Exception:
                    pass
    print(f"Resuming: {len(done)} already done" if done else "Fresh run")

    s3 = boto3.client("s3", region_name=args.region)
    manifest = manifest_path.open("a")
    max_bytes = int(args.max_mb * 1024 * 1024)
    st = {
        "seen": 0,
        "extracted": 0,
        "skip_type": 0,
        "skip_scanned": 0,
        "skip_big": 0,
        "skip_empty": 0,
        "errors": 0,
        "poison": 0,
        "bytes": 0,
    }
    t0 = time.time()

    # Crash recovery. Native libs (PyMuPDF) can SEGFAULT on a malformed file, which
    # Python cannot catch. We write a breadcrumb (the key we're about to touch) before
    # each risky op; if the process dies, the next run finds the breadcrumb and
    # quarantines that key so it's skipped instead of crashing again. `.cursor` lets a
    # restart resume listing right where it left off instead of re-scanning ~1M keys.
    cur_path = out_root / ".current"
    cursor_path = out_root / ".cursor"
    quar_path = out_root / "quarantine.txt"
    quarantine: set[str] = set()
    if quar_path.exists():
        quarantine.update(
            x.strip() for x in quar_path.read_text().splitlines() if x.strip()
        )
    if cur_path.exists():
        poisoned = cur_path.read_text().strip()
        if poisoned and poisoned not in quarantine:
            quarantine.add(poisoned)
            with quar_path.open("a") as q:
                q.write(poisoned + "\n")
            print(f"QUARANTINED (crashed last run): {poisoned}")
        cur_path.write_text("")
    start_after = cursor_path.read_text().strip() if cursor_path.exists() else ""
    if start_after:
        print(f"Resuming listing after: {start_after}")

    try:
        for key, size in iter_keys(s3, args.bucket, args.prefix, start_after):
            ext = file_ext(key)
            if ext not in HANDLED:
                st["skip_type"] += 1
                continue
            if key in done:
                continue
            if key in quarantine:
                st["poison"] += 1
                continue
            st["seen"] += 1
            if size and size > max_bytes:
                st["skip_big"] += 1
                cursor_path.write_text(key)
                continue
            cur_path.write_text(key)  # breadcrumb — see crash recovery above
            try:
                try:
                    data = s3.get_object(Bucket=args.bucket, Key=key)["Body"].read()
                    st["bytes"] += len(data)
                    text = extract(data, ext).strip()
                except _Skip as s:
                    if str(s) == "pdf-scanned":
                        st["skip_scanned"] += 1
                    else:
                        st["skip_type"] += 1
                    continue
                except Exception as exc:  # noqa: BLE001
                    st["errors"] += 1
                    sys.stderr.write(f"  ! {key}: {exc}\n")
                    continue

                if not text:
                    st["skip_empty"] += 1
                    continue

                fname = key.rsplit("/", 1)[-1]
                stem = safe(f"{job_id_from_key(key)}__{fname}")
                (docs_dir / f"{stem}.txt").write_text(text, encoding="utf-8")
                meta = {
                    "metadataAttributes": {
                        # tenant_id + kb_id are REQUIRED: Numa filters retrieval on
                        # andAll[tenant_id==CLIENT_NAME, kb_id==folder]. Without them
                        # chat returns "No results" even though docs are vectorised.
                        "tenant_id": args.tenant_id,
                        "kb_id": args.kb_id,
                        "uploader_id": "system",
                        "source_bucket": args.bucket,
                        "source_key": key,
                        "job_id": job_id_from_key(key),
                        "file_name": fname,
                    }
                }
                (docs_dir / f"{stem}.txt.metadata.json").write_text(
                    json.dumps(meta, ensure_ascii=False), encoding="utf-8"
                )
                manifest.write(json.dumps({"key": key, "chars": len(text)}) + "\n")
                manifest.flush()
                st["extracted"] += 1

                if st["extracted"] % 100 == 0:
                    _progress(st, t0)
                if args.max_files and st["extracted"] >= args.max_files:
                    break
            finally:
                cursor_path.write_text(key)
                cur_path.write_text("")
    except KeyboardInterrupt:
        print("\nInterrupted — manifest saved, safe to re-run.")
    finally:
        manifest.close()

    _progress(st, t0, final=True)
    print(f"\nCorpus: {docs_dir}")
    print("Next: aws s3 sync", str(docs_dir), "s3://<kb-bucket>/synergy/")


def _progress(st: dict, t0: float, final: bool = False) -> None:
    mins = (time.time() - t0) / 60
    gb = st["bytes"] / (1024**3)
    print(
        f"[{'DONE' if final else '...'}] extracted={st['extracted']} "
        f"skip(type={st['skip_type']},scanned={st['skip_scanned']},"
        f"big={st['skip_big']},empty={st['skip_empty']},poison={st['poison']}) "
        f"err={st['errors']} dl={gb:.2f}GB {mins:.1f}min",
        flush=True,
    )


if __name__ == "__main__":
    main()
